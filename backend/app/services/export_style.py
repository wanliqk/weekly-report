"""Visual styling for the daily-report export workbook.

Keeps every openpyxl style/layout concern (title banner, header shading,
borders, column width, wrap text, row height, frozen header) in one place
so `export.py` only ever deals with *what* data goes in which cell, never
*how* it looks. Nothing here reads or writes business data.
"""

import unicodedata
from dataclasses import dataclass

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

REPORT_TITLE = "每日工作汇报"

_FONT_NAME = "微软雅黑"
_BRAND_COLOR = "FF1F3864"

_TITLE_FONT = Font(name=_FONT_NAME, size=18, bold=True, color=_BRAND_COLOR)
_TITLE_ALIGNMENT = Alignment(horizontal="center", vertical="center")
_TITLE_ROW_HEIGHT = 34.0

_HEADER_FONT = Font(name=_FONT_NAME, size=11, bold=True, color="FFFFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor=_BRAND_COLOR)
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_HEADER_ROW_HEIGHT = 26.0

_BODY_FONT = Font(name=_FONT_NAME, size=10.5, color="FF333333")
_BODY_ALIGNMENT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_BODY_ROW_MIN_HEIGHT = 20.0
_BODY_ROW_LINE_HEIGHT = 15.0

_BLOCK_TITLE_FONT = Font(name=_FONT_NAME, size=11, bold=True, color="FF1F3864")
_BLOCK_TITLE_ALIGNMENT = Alignment(horizontal="left", vertical="center")
_BLOCK_TITLE_ROW_HEIGHT = 22.0

_BORDER_COLOR = "FFBFBFBF"
_CELL_BORDER = Border(
    left=Side(style="thin", color=_BORDER_COLOR),
    right=Side(style="thin", color=_BORDER_COLOR),
    top=Side(style="thin", color=_BORDER_COLOR),
    bottom=Side(style="thin", color=_BORDER_COLOR),
)

_MIN_COLUMN_WIDTH = 12.0
_MAX_COLUMN_WIDTH = 42.0
_COLUMN_WIDTH_PADDING = 4.0


@dataclass(frozen=True)
class ProjectListBlockLayout:
    """Row span of one `PROJECT_LIST` sub-table appended below the main table.

    One block is one (day, field) pair with at least one entry: a bold
    title row identifying the day/owner/field, a 3-column header row
    (`项目`/`工作内容`/`完成状态`), and its data rows.
    """

    title_row: int
    header_row: int
    first_data_row: int
    last_data_row: int
    column_count: int = 3


def style_report_sheet(
    sheet: Worksheet,
    *,
    header_row: int,
    first_data_row: int,
    last_data_row: int,
    column_count: int,
    project_list_blocks: tuple[ProjectListBlockLayout, ...] = (),
) -> None:
    """Applies the enterprise daily-report look to an already-populated sheet.

    Expects an empty row directly above `header_row` for the title, headers
    already written on `header_row`, and data rows `first_data_row..
    last_data_row` already filled across columns `1..column_count`. Only
    ever sets formatting (and the title cell's text) — never touches the
    data cell values.

    `project_list_blocks` styles each appended `PROJECT_LIST` sub-table the
    same way (bold header, borders, wrap text) and folds their cells into
    the same column-width pass as the main table, since a block's 3 columns
    reuse the sheet's leftmost column letters and must not be narrower than
    whatever the main table already needs there.
    """
    _style_title(sheet, row=header_row - 1, column_count=column_count)
    _style_header(sheet, row=header_row, column_count=column_count)
    _style_body(sheet, first_row=first_data_row, last_row=last_data_row, column_count=column_count)
    for block in project_list_blocks:
        _style_block_title(sheet, row=block.title_row)
        _style_header(sheet, row=block.header_row, column_count=block.column_count)
        _style_body(
            sheet,
            first_row=block.first_data_row,
            last_row=block.last_data_row,
            column_count=block.column_count,
        )
    block_last_rows = [block.last_data_row for block in project_list_blocks]
    block_column_counts = [block.column_count for block in project_list_blocks]
    _autosize_columns(
        sheet,
        header_row=header_row,
        first_data_row=first_data_row,
        last_data_row=max([last_data_row, *block_last_rows]),
        column_count=max([column_count, *block_column_counts]),
    )
    sheet.freeze_panes = sheet.cell(row=first_data_row, column=1).coordinate


def _style_block_title(sheet: Worksheet, *, row: int) -> None:
    cell = sheet.cell(row=row, column=1)
    cell.font = _BLOCK_TITLE_FONT
    cell.alignment = _BLOCK_TITLE_ALIGNMENT
    sheet.row_dimensions[row].height = _BLOCK_TITLE_ROW_HEIGHT


def _style_title(sheet: Worksheet, *, row: int, column_count: int) -> None:
    cell = sheet.cell(row=row, column=1)
    cell.value = REPORT_TITLE
    cell.font = _TITLE_FONT
    cell.alignment = _TITLE_ALIGNMENT
    if column_count > 1:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=column_count)
    sheet.row_dimensions[row].height = _TITLE_ROW_HEIGHT


def _style_header(sheet: Worksheet, *, row: int, column_count: int) -> None:
    for column in range(1, column_count + 1):
        cell = sheet.cell(row=row, column=column)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _CELL_BORDER
    sheet.row_dimensions[row].height = _HEADER_ROW_HEIGHT


def _style_body(sheet: Worksheet, *, first_row: int, last_row: int, column_count: int) -> None:
    for row in range(first_row, last_row + 1):
        max_lines = 1
        for column in range(1, column_count + 1):
            cell = sheet.cell(row=row, column=column)
            cell.font = _BODY_FONT
            cell.alignment = _BODY_ALIGNMENT
            cell.border = _CELL_BORDER
            if isinstance(cell.value, str):
                max_lines = max(max_lines, cell.value.count("\n") + 1)
        sheet.row_dimensions[row].height = max(
            _BODY_ROW_MIN_HEIGHT, max_lines * _BODY_ROW_LINE_HEIGHT
        )


def _autosize_columns(
    sheet: Worksheet,
    *,
    header_row: int,
    first_data_row: int,
    last_data_row: int,
    column_count: int,
) -> None:
    for column in range(1, column_count + 1):
        widest = _display_width(str(sheet.cell(row=header_row, column=column).value or ""))
        for row in range(first_data_row, last_data_row + 1):
            widest = max(widest, *_cell_line_widths(sheet.cell(row=row, column=column).value))
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = min(
            _MAX_COLUMN_WIDTH, max(_MIN_COLUMN_WIDTH, widest + _COLUMN_WIDTH_PADDING)
        )


def _display_width(text: str) -> float:
    """Approximates a string's on-screen width in Excel "character" units.

    Excel sizes columns in units close to one ASCII digit's width; CJK
    characters render roughly twice as wide, so plain `len()` badly
    under-estimates the width of Chinese report content and truncates it.
    """
    width = 0.0
    for char in text:
        width += 2.0 if unicodedata.east_asian_width(char) in "WF" else 1.0
    return width


def _cell_line_widths(value: object) -> list[float]:
    if value is None:
        return [0.0]
    lines = str(value).splitlines()
    return [_display_width(line) for line in lines] or [0.0]
