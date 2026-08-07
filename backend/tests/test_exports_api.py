from io import BytesIO
from typing import Any, cast

import openpyxl
import pytest
from conftest import STAGE5_RUNTIME_HEADERS
from fastapi.testclient import TestClient

from app.core.config import Settings
from app.core.ulid import generate_ulid


def _create(client: TestClient, headers: dict[str, str], work_date: str) -> dict[str, Any]:
    response = client.post(
        "/api/v1/daily-reports",
        headers=headers,
        json={"work_date": work_date, "client_request_id": generate_ulid()},
    )
    assert response.status_code == 200, response.text
    return cast(dict[str, Any], response.json()["data"])


def _default_content(report: dict[str, Any]) -> dict[str, object]:
    return {
        field["field_key"]: f"{field['label']}内容"
        for field in report["template_snapshot"]
        if field["enabled"] and field["required"]
    }


def _archive(client: TestClient, headers: dict[str, str], work_date: str) -> dict[str, Any]:
    report = _create(client, headers, work_date)
    save = client.patch(
        f"/api/v1/daily-reports/{report['id']}",
        headers=headers,
        json={"version": report["version"], "content": _default_content(report)},
    )
    assert save.status_code == 200, save.text
    submit = client.post(
        f"/api/v1/daily-reports/{report['id']}/submit",
        headers=headers,
        json={"version": save.json()["data"]["version"]},
    )
    assert submit.status_code == 200, submit.text
    archived_day = client.post(
        f"/api/v1/daily-report-days/{work_date}/archive",
        headers=headers,
        json={"confirm_archive": True},
    )
    assert archived_day.status_code == 200, archived_day.text
    entry = client.get(f"/api/v1/daily-reports/{report['id']}", headers=headers)
    assert entry.status_code == 200, entry.text
    return cast(dict[str, Any], entry.json()["data"])


def _create_user_headers(client: TestClient, admin_headers: dict[str, str]) -> dict[str, str]:
    create = client.post(
        "/api/v1/users",
        headers=admin_headers,
        json={
            "username": "alice",
            "password": "alice secure password",
            "display_name": "Alice",
            "role": "user",
        },
    )
    assert create.status_code == 200, create.text
    login = client.post(
        "/api/v1/auth/login",
        headers=STAGE5_RUNTIME_HEADERS,
        json={"username": "alice", "password": "alice secure password"},
    )
    assert login.status_code == 200, login.text
    token = cast(str, login.json()["data"]["access_token"])
    return {**STAGE5_RUNTIME_HEADERS, "Authorization": f"Bearer {token}"}


def test_export_by_ids_returns_succeeded_job_with_downloadable_xlsx(
    stage5_context: tuple[TestClient, dict[str, str], Settings],
) -> None:
    client, headers, _settings = stage5_context
    archived = _archive(client, headers, "2026-08-01")

    created = client.post(
        "/api/v1/daily-report-exports",
        headers=headers,
        json={"report_ids": [archived["id"]]},
    )
    assert created.status_code == 200, created.text
    job = created.json()["data"]
    assert job["status"] == "succeeded"
    assert job["record_count"] == 1
    assert job["file_name"].endswith(".xlsx")

    fetched = client.get(f"/api/v1/daily-report-exports/{job['id']}", headers=headers)
    assert fetched.status_code == 200
    assert fetched.json()["data"]["status"] == "succeeded"

    downloaded = client.get(f"/api/v1/daily-report-exports/{job['id']}/file", headers=headers)
    assert downloaded.status_code == 200
    assert downloaded.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert job["file_name"] in downloaded.headers["content-disposition"]
    workbook = openpyxl.load_workbook(BytesIO(downloaded.content))
    sheet = workbook.active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:3] == ("工作日期", "提交时间", "归档时间")
    assert rows[1][0] == "2026-08-01"


def test_export_file_exposes_content_disposition_to_cross_origin_renderer(
    stage5_context: tuple[TestClient, dict[str, str], Settings],
) -> None:
    # Renderer downloads are cross-origin fetches (127.0.0.1:<port> vs the
    # Vite origin); Content-Disposition is not on the browser's default CORS
    # response-header safelist, so without `expose_headers` the renderer can
    # never read the server-suggested file name (regression: real Electron
    # smoke test showed the save dialog silently falling back to a generic
    # default name).
    client, headers, _settings = stage5_context
    archived = _archive(client, headers, "2026-08-01")
    created = client.post(
        "/api/v1/daily-report-exports",
        headers=headers,
        json={"report_ids": [archived["id"]]},
    )
    job_id = created.json()["data"]["id"]

    downloaded = client.get(
        f"/api/v1/daily-report-exports/{job_id}/file",
        headers={**headers, "Origin": "http://127.0.0.1:5173"},
    )

    assert downloaded.status_code == 200
    assert "content-disposition" in downloaded.headers["access-control-expose-headers"].lower()


@pytest.mark.parametrize(
    "payload",
    [
        {},
        {"report_ids": ["01AAAAAAAAAAAAAAAAAAAAAAAA"], "filter": {}},
    ],
)
def test_export_requires_exactly_one_selection_mode(
    stage5_context: tuple[TestClient, dict[str, str], Settings],
    payload: dict[str, Any],
) -> None:
    client, headers, _settings = stage5_context
    response = client.post("/api/v1/daily-report-exports", headers=headers, json=payload)
    assert response.status_code == 400
    assert response.json()["code"] == 40001


def test_export_rejects_selection_mixing_archived_and_non_archived_ids(
    stage5_context: tuple[TestClient, dict[str, str], Settings],
) -> None:
    client, headers, _settings = stage5_context
    archived = _archive(client, headers, "2026-08-01")
    draft = _create(client, headers, "2026-08-02")

    response = client.post(
        "/api/v1/daily-report-exports",
        headers=headers,
        json={"report_ids": [archived["id"], draft["id"]]},
    )

    assert response.status_code == 400
    assert response.json()["code"] == 40001
    assert response.json()["data"]["invalid_report_ids"] == [draft["id"]]


def test_export_by_filter_only_includes_archived_reports_in_range(
    stage5_context: tuple[TestClient, dict[str, str], Settings],
) -> None:
    client, headers, _settings = stage5_context
    archived = _archive(client, headers, "2026-08-02")
    _create(client, headers, "2026-08-03")

    response = client.post(
        "/api/v1/daily-report-exports",
        headers=headers,
        json={"filter": {"date_from": "2026-08-01", "date_to": "2026-08-31"}},
    )

    assert response.status_code == 200, response.text
    job = response.json()["data"]
    assert job["record_count"] == 1
    downloaded = client.get(f"/api/v1/daily-report-exports/{job['id']}/file", headers=headers)
    workbook = openpyxl.load_workbook(BytesIO(downloaded.content))
    sheet = workbook.active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[1][0] == archived["work_date"]


def test_export_job_and_file_are_isolated_per_owner(
    stage5_context: tuple[TestClient, dict[str, str], Settings],
) -> None:
    client, admin_headers, _settings = stage5_context
    archived = _archive(client, admin_headers, "2026-08-01")
    created = client.post(
        "/api/v1/daily-report-exports",
        headers=admin_headers,
        json={"report_ids": [archived["id"]]},
    )
    job_id = created.json()["data"]["id"]
    alice_headers = _create_user_headers(client, admin_headers)

    foreign_selection = client.post(
        "/api/v1/daily-report-exports",
        headers=alice_headers,
        json={"report_ids": [archived["id"]]},
    )
    get_job = client.get(f"/api/v1/daily-report-exports/{job_id}", headers=alice_headers)
    download = client.get(f"/api/v1/daily-report-exports/{job_id}/file", headers=alice_headers)

    assert foreign_selection.status_code == 400
    assert foreign_selection.json()["data"]["invalid_report_ids"] == [archived["id"]]
    assert get_job.status_code == 404
    assert get_job.json()["code"] == 40401
    assert download.status_code == 404
    assert download.json()["code"] == 40401


def test_export_routes_require_authentication(
    stage5_context: tuple[TestClient, dict[str, str], Settings],
) -> None:
    client, _headers, _settings = stage5_context
    response = client.post(
        "/api/v1/daily-report-exports",
        headers={"X-Runtime-Secret": STAGE5_RUNTIME_HEADERS["X-Runtime-Secret"]},
        json={"filter": {}},
    )
    assert response.status_code == 401
    assert response.json()["code"] == 40102
