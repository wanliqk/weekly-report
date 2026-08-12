import json
from collections.abc import AsyncIterator
from datetime import UTC, date, datetime
from pathlib import Path

import openpyxl
import pytest
from conftest import STAGE5_PASSWORD
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.ext.asyncio import AsyncEngine, AsyncSession

from app.core.config import Settings
from app.core.paths import ensure_runtime_directories
from app.core.ulid import generate_ulid
from app.db.engine import create_engine
from app.db.migrate import run_startup_migrations
from app.db.session import create_session_factory
from app.models import DailyReportDay, ExportJob, User
from app.repositories.daily_report_day import DailyReportDayRepository
from app.repositories.template import TemplateRepository
from app.schemas.export import ExportFilter
from app.services.bootstrap import BootstrapService
from app.services.export import (
    EXPORT_MEDIA_TYPE,
    ExportFileNotAvailableError,
    ExportSelectionInvalidError,
    ExportService,
    _safe_filename_component,
)
from app.services.export_style import REPORT_TITLE

_FIXED_NOW = datetime(2026, 8, 5, 12, 0, 0, tzinfo=UTC)
_DATE_HEADER = "日期"


def _past_clock() -> datetime:
    """A fixed instant far enough in the past that `expires_at` (+24h) is
    already behind any real wall-clock `now`, regardless of when tests run.
    """
    return datetime(2000, 1, 1, tzinfo=UTC)


@pytest.fixture
def export_settings(tmp_path: Path) -> Settings:
    settings = Settings(
        environment="test",
        data_dir=tmp_path / "data",
        backup_dir=tmp_path / "backups",
        export_temp_dir=tmp_path / "exports",
    )
    ensure_runtime_directories(settings)
    run_startup_migrations(settings)
    return settings


@pytest.fixture
async def export_engine(export_settings: Settings) -> AsyncIterator[AsyncEngine]:
    engine = create_engine(export_settings)
    try:
        yield engine
    finally:
        await engine.dispose()


def _fields(pairs: list[tuple[str, str]]) -> list[dict[str, object]]:
    return [
        {
            "field_key": key,
            "label": label,
            "description": "",
            "field_type": "text",
            "required": False,
            "enabled": True,
            "sort_order": index * 10,
            "options": [],
            "core_type": None,
        }
        for index, (key, label) in enumerate(pairs)
    ]


async def _default_template_version_id(session: AsyncSession, owner_id: str) -> str:
    repository = TemplateRepository(session)
    template = await repository.get_for_owner(owner_id)
    assert template is not None
    version = await repository.get_version(template.id, template.current_version_no)
    assert version is not None
    return version.id


def _snapshot(
    *, work_date: date, entries: list[tuple[str, list[dict[str, object]], dict[str, object]]]
) -> str:
    return json.dumps(
        {
            "schema_version": 2,
            "work_date": work_date.isoformat(),
            "entries": [
                {
                    "daily_report_id": entry_id,
                    "submitted_at": _FIXED_NOW.isoformat(),
                    "template_version_id": "version",
                    "template_snapshot": fields,
                    "content": content,
                }
                for entry_id, fields, content in entries
            ],
        },
        ensure_ascii=False,
    )


async def _archived_day(
    session: AsyncSession,
    *,
    owner_id: str,
    work_date: date,
    fields: list[dict[str, object]],
    content: dict[str, object],
) -> DailyReportDay:
    return await _archived_day_multi(
        session, owner_id=owner_id, work_date=work_date, entries=[(fields, content)]
    )


async def _archived_day_multi(
    session: AsyncSession,
    *,
    owner_id: str,
    work_date: date,
    entries: list[tuple[list[dict[str, object]], dict[str, object]]],
) -> DailyReportDay:
    await _default_template_version_id(session, owner_id)  # ensures template exists
    entry_ids = [generate_ulid() for _ in entries]
    day = DailyReportDay(
        id=generate_ulid(),
        user_id=owner_id,
        work_date=work_date,
        status="archived",
        archive_snapshot_json=_snapshot(
            work_date=work_date,
            entries=[
                (entry_id, fields, content)
                for entry_id, (fields, content) in zip(entry_ids, entries, strict=True)
            ],
        ),
        source_count=len(entries),
        archived_by=owner_id,
        archived_at=_FIXED_NOW,
        version=2,
    )
    day_repository = DailyReportDayRepository(session)
    await day_repository.add(day)
    await session.commit()
    return day


async def _draft_day(session: AsyncSession, *, owner_id: str, work_date: date) -> DailyReportDay:
    day = DailyReportDay(
        id=generate_ulid(), user_id=owner_id, work_date=work_date, status="open", version=1
    )
    await DailyReportDayRepository(session).add(day)
    await session.commit()
    return day


async def _bootstrap_user(session: AsyncSession) -> User:
    return await BootstrapService(session).bootstrap(
        username="owner", password=STAGE5_PASSWORD, display_name="Owner"
    )


def _workbook_path(job: ExportJob) -> Path:
    assert job.file_path is not None
    return Path(job.file_path)


def _row_values(sheet: Worksheet, row: int) -> list[object]:
    """Reads every one of a row's cells across the sheet's full column count.

    Deliberately does *not* trim trailing `None`s: in the merged-cell layout
    a trailing `None` can be real data (a day-level column left blank on a
    project row it was merged away from, or a `PROJECT_LIST` column with no
    entries for that day) — callers must spell out the exact expected width.
    """
    return [cell.value for cell in sheet[row]]


async def test_create_from_ids_merges_columns_across_snapshots_and_marks_succeeded(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)
        older = await _archived_day(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 3),
            fields=_fields([("k1", "今日进展"), ("k2", "明日计划")]),
            content={"k1": "写文档", "k2": "写测试"},
        )
        newer = await _archived_day(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 4),
            fields=_fields([("k1", "今日工作内容"), ("k2", "明日计划"), ("k3", "风险")]),
            content={"k1": "评审代码", "k2": "发布", "k3": "无"},
        )

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=[newer.id, older.id], filter_=None
        )

    assert job.status == "succeeded"
    assert job.record_count == 2
    assert job.file_name == "日报_Owner_2026年8月4日.xlsx"
    workbook_path = _workbook_path(job)
    assert workbook_path.parent == export_settings.export_temp_dir.resolve()

    workbook = openpyxl.load_workbook(workbook_path)
    sheet = workbook.active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == REPORT_TITLE
    assert rows[1] == (_DATE_HEADER, "今日工作内容", "明日计划", "风险")
    assert rows[2][0] == "2026-08-03"
    assert rows[2][1:] == ("写文档", "写测试", None)
    assert rows[3][0] == "2026-08-04"
    assert rows[3][1:] == ("评审代码", "发布", "无")


async def test_create_merges_multiple_source_entries_for_one_day_into_numbered_lines(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    """`docs/方案设计.md` §9.2: multi-source values render as `[1] .../[2] ...`
    lines within a single cell, and a single-source cell keeps its plain value.
    """
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)
        day = await _archived_day_multi(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 5),
            entries=[
                (_fields([("k1", "今日工作内容")]), {"k1": "上午写文档"}),
                (_fields([("k1", "今日工作内容")]), {"k1": "下午写测试"}),
            ],
        )

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=[day.id], filter_=None
        )

    workbook = openpyxl.load_workbook(_workbook_path(job))
    sheet = workbook.active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[2][1] == "[1] 上午写文档\n[2] 下午写测试"


def _field(
    key: str, label: str, *, sort_order: int = 0, show_in_export: bool = True
) -> dict[str, object]:
    return {
        "field_key": key,
        "label": label,
        "description": "",
        "field_type": "text",
        "required": False,
        "enabled": True,
        "show_in_export": show_in_export,
        "sort_order": sort_order,
        "options": [],
        "core_type": None,
    }


async def test_create_excludes_a_field_marked_hidden_from_export(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    """A field with `show_in_export=False` contributes no column at all — its

    value is never written for any day, and it does not count toward the
    same-label disambiguation applied to the columns that remain visible.
    """
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)
        await _archived_day(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 3),
            fields=[
                _field("k1", "备注", sort_order=0, show_in_export=False),
                _field("k2", "今日工作内容", sort_order=10),
            ],
            content={"k1": "仅供内部台账使用", "k2": "写文档"},
        )

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=None, filter_=ExportFilter()
        )

    workbook = openpyxl.load_workbook(_workbook_path(job))
    sheet = workbook.active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[1] == (_DATE_HEADER, "今日工作内容")
    assert rows[2] == ("2026-08-03", "写文档")


async def test_create_applies_most_recent_snapshot_wins_to_show_in_export(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    """Mirrors the existing "most recent label wins" merge rule

    (`test_create_from_ids_merges_columns_across_snapshots_and_marks_succeeded`):
    a field later hidden from export drops out of the merged column list even
    though an older day's snapshot still had it visible.
    """
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)
        await _archived_day(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 3),
            fields=[_field("k1", "今日工作内容", show_in_export=True)],
            content={"k1": "写文档"},
        )
        await _archived_day(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 4),
            fields=[_field("k1", "今日工作内容", show_in_export=False)],
            content={"k1": "评审代码"},
        )

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=None, filter_=ExportFilter()
        )

    workbook = openpyxl.load_workbook(_workbook_path(job))
    sheet = workbook.active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[1] == (_DATE_HEADER,)
    assert rows[2] == ("2026-08-03",)
    assert rows[3] == ("2026-08-04",)


def _project_list_field(key: str, label: str) -> dict[str, object]:
    return {
        "field_key": key,
        "label": label,
        "description": "",
        "field_type": "PROJECT_LIST",
        "required": False,
        "enabled": True,
        "sort_order": 0,
        "options": [],
        "core_type": None,
    }


def _project_list_and_plan_fields(
    project_key: str = "k1", plan_key: str = "k2"
) -> list[dict[str, object]]:
    """A `PROJECT_LIST` field followed by a plain field, matching the

    `日期/类别/工作项目/工作步骤/权重/预计完成时间节点/实际完成时间/责任人/
    协助人/所需资源支持/实际完成情况及解决措施/明日工作计划` column order.
    PROD-030 expands `PROJECT_LIST` to 10 sub-columns (`责任人` remains a
    project-level value; the old base owner column was removed by PROD-027).
    """
    return [
        _project_list_field(project_key, "今日工作"),
        {
            "field_key": plan_key,
            "label": "明日工作计划",
            "description": "",
            "field_type": "textarea",
            "required": False,
            "enabled": True,
            "sort_order": 10,
            "options": [],
            "core_type": None,
        },
    ]


def _merge_ranges(sheet: Worksheet) -> set[str]:
    return {str(cell_range) for cell_range in sheet.merged_cells.ranges}


async def test_create_merges_day_level_columns_across_several_project_rows(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    """`ai-docs/decisions.md` PROD-024: a `PROJECT_LIST` field stays inline in

    the main table (not a separate block) as its 10 sub-columns (PROD-030)
    at its own sort position; every other field (`日期`/plain fields such as
    `明日工作计划`) is written once and vertically merged across the day's
    project rows so it is never repeated.
    """
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)
        await _archived_day(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 3),
            fields=_project_list_and_plan_fields(),
            content={
                "k1": [
                    {"project": "测试项目1", "content": "测试测试"},
                    {"project": "测试项目2", "content": "测试测试"},
                ],
                "k2": "明天计划",
            },
        )

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=None, filter_=ExportFilter()
        )

    workbook = openpyxl.load_workbook(_workbook_path(job))
    sheet = workbook.active
    assert sheet is not None
    assert _row_values(sheet, 2) == [
        "日期",
        "类别",
        "工作项目",
        "工作步骤",
        "权重",
        "预计完成时间节点",
        "实际完成时间",
        "责任人",
        "协助人",
        "所需资源支持",
        "实际完成情况及解决措施",
        "明日工作计划",
    ]
    assert _row_values(sheet, 3) == [
        "2026-08-03",
        "重要",
        "测试项目1",
        "测试测试",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        "明天计划",
    ]
    assert _row_values(sheet, 4) == [
        None,
        "重要",
        "测试项目2",
        "测试测试",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
    ]
    assert sheet.max_row == 4

    merges = _merge_ranges(sheet)
    assert "A3:A4" in merges  # 日期
    assert "L3:L4" in merges  # 明日工作计划
    # PROJECT_LIST 的 10 个子列各自随条目变化 必须都不合并。
    assert not any(
        cell_range.startswith(("B3", "C3", "D3", "E3", "F3", "G3", "H3", "I3", "J3", "K3"))
        for cell_range in merges
    )


async def test_create_handles_a_single_project_without_merging(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    """Requirement 6: exactly one project needs no merge — just one plain row."""
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)
        await _archived_day(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 3),
            fields=[_project_list_field("k1", "今日工作")],
            content={"k1": [{"project": "测试项目1", "content": "测试测试"}]},
        )

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=None, filter_=ExportFilter()
        )

    workbook = openpyxl.load_workbook(_workbook_path(job))
    sheet = workbook.active
    assert sheet is not None
    assert _row_values(sheet, 3) == [
        "2026-08-03",
        "重要",
        "测试项目1",
        "测试测试",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
    ]
    assert sheet.max_row == 3
    # Only the title-row merge exists; no day-level merge was needed for one row.
    assert _merge_ranges(sheet) == {"A1:K1"}


async def test_create_handles_an_empty_project_list(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    """Requirement 7: zero projects still produces the day's row, with blank

    sub-columns and (with only one row) no merge.
    """
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)
        await _archived_day(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 3),
            fields=[_project_list_field("k1", "今日工作")],
            content={"k1": []},
        )

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=None, filter_=ExportFilter()
        )

    workbook = openpyxl.load_workbook(_workbook_path(job))
    sheet = workbook.active
    assert sheet is not None
    assert _row_values(sheet, 3) == [
        "2026-08-03",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
    ]
    assert sheet.max_row == 3
    assert _merge_ranges(sheet) == {"A1:K1"}


async def test_create_flattens_project_list_entries_from_every_source_without_reordering(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    """A day archived from several source entries still produces one row per

    project, in submission order — sources are not distinguished with `[N]`
    numbering the way single-value fields are.
    """
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)
        day = await _archived_day_multi(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 5),
            entries=[
                (
                    [_project_list_field("k1", "今日工作")],
                    {"k1": [{"project": "项目A", "content": "上午任务"}]},
                ),
                (
                    [_project_list_field("k1", "今日工作")],
                    {"k1": [{"project": "项目B", "content": "下午任务"}]},
                ),
            ],
        )

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=[day.id], filter_=None
        )

    workbook = openpyxl.load_workbook(_workbook_path(job))
    sheet = workbook.active
    assert sheet is not None
    assert _row_values(sheet, 3) == [
        "2026-08-05",
        "重要",
        "项目A",
        "上午任务",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
    ]
    assert _row_values(sheet, 4) == [
        None,
        "重要",
        "项目B",
        "下午任务",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
    ]
    assert "A3:A4" in _merge_ranges(sheet)


async def test_create_defuses_formula_like_project_and_content_cells(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    """Every `PROJECT_LIST` sub-column is a standalone free-text cell (there is

    no fixed-enum column left, unlike pre-PROD-028's `进度`), so each is
    independently at risk of Excel formula promotion and must go through
    `_defuse_formula` on its own; spot-checked here on the two new fields,
    `project`/`content`, and `owner` across `_PROJECT_LIST_FIELD_ORDER`.
    """
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)
        await _archived_day(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 5),
            fields=[_project_list_field("k1", "今日工作")],
            content={
                "k1": [
                    {
                        "category": "=CATEGORY()",
                        "project": '=HYPERLINK("http://evil.example","x")',
                        "content": "=cmd|' /C calc'!A0",
                        "weight": "=30/100",
                        "owner": "=SUM(A1:A2)",
                    }
                ]
            },
        )

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=None, filter_=ExportFilter()
        )

    workbook = openpyxl.load_workbook(_workbook_path(job))
    sheet = workbook.active
    assert sheet is not None
    assert _row_values(sheet, 3) == [
        "2026-08-05",
        "'=CATEGORY()",
        '\'=HYPERLINK("http://evil.example","x")',
        "'=cmd|' /C calc'!A0",
        "'=30/100",
        None,
        None,
        "'=SUM(A1:A2)",
        None,
        None,
        None,
    ]
    assert sheet.cell(row=3, column=2).data_type != "f"
    assert sheet.cell(row=3, column=3).data_type != "f"
    assert sheet.cell(row=3, column=4).data_type != "f"
    assert sheet.cell(row=3, column=5).data_type != "f"
    assert sheet.cell(row=3, column=8).data_type != "f"


async def test_create_merges_independently_per_day(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    """Two days with different project counts each get their own correct

    row span and merge range; one day's expansion must not affect another's.
    """
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)
        await _archived_day(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 4),
            fields=[_project_list_field("k1", "今日工作")],
            content={"k1": [{"project": "项目A", "content": "内容A"}]},
        )
        await _archived_day(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 5),
            fields=[_project_list_field("k1", "今日工作")],
            content={
                "k1": [
                    {"project": "项目B1", "content": "内容B1"},
                    {"project": "项目B2", "content": "内容B2"},
                ]
            },
        )

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=None, filter_=ExportFilter()
        )

    workbook = openpyxl.load_workbook(_workbook_path(job))
    sheet = workbook.active
    assert sheet is not None
    assert _row_values(sheet, 3) == [
        "2026-08-04",
        "重要",
        "项目A",
        "内容A",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
    ]
    assert _row_values(sheet, 4) == [
        "2026-08-05",
        "重要",
        "项目B1",
        "内容B1",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
    ]
    assert _row_values(sheet, 5) == [
        None,
        "重要",
        "项目B2",
        "内容B2",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
    ]
    assert sheet.max_row == 5
    merges = _merge_ranges(sheet)
    assert "A4:A5" in merges
    assert not any(cell_range.startswith("A3") for cell_range in merges)


async def test_create_defuses_formula_like_content_but_preserves_bullet_dashes(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)
        await _archived_day(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 5),
            fields=_fields([("k1", "今日工作内容")]),
            content={"k1": '=HYPERLINK("http://evil.example","x")'},
        )
        await _archived_day(
            session,
            owner_id=user.id,
            work_date=date(2026, 8, 6),
            fields=_fields([("k1", "今日工作内容")]),
            content={"k1": "-完成需求分析\n-编写代码"},
        )

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=None, filter_=ExportFilter()
        )

    workbook = openpyxl.load_workbook(_workbook_path(job))
    sheet = workbook.active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    formula_cell, bullet_cell = rows[2][1], rows[3][1]
    assert formula_cell == '\'=HYPERLINK("http://evil.example","x")'
    assert bullet_cell == "-完成需求分析\n-编写代码"

    formula_row = next(row for row in sheet.iter_rows() if row[0].value == "2026-08-05")
    assert formula_row[1].data_type != "f"


async def test_create_rejects_selection_containing_foreign_or_non_archived_ids(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)
        archived = await _archived_day(
            session, owner_id=user.id, work_date=date(2026, 8, 5), fields=[], content={}
        )
        draft = await _draft_day(session, owner_id=user.id, work_date=date(2026, 8, 6))

    async with session_factory() as session:
        with pytest.raises(ExportSelectionInvalidError) as excinfo:
            await ExportService(session, export_settings).create(
                user.id,
                daily_report_day_ids=[archived.id, draft.id, "01MISSINGMISSINGMISSINGMI"],
                filter_=None,
            )

    assert set(excinfo.value.data["invalid_daily_report_day_ids"]) == {
        draft.id,
        "01MISSINGMISSINGMISSINGMI",
    }


async def test_create_from_filter_only_includes_archived_reports_in_range(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)
        await _archived_day(
            session,
            owner_id=user.id,
            work_date=date(2026, 7, 31),
            fields=[],
            content={},
        )
        in_range = await _archived_day(
            session, owner_id=user.id, work_date=date(2026, 8, 2), fields=[], content={}
        )
        await _draft_day(session, owner_id=user.id, work_date=date(2026, 8, 3))

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id,
            daily_report_day_ids=None,
            filter_=ExportFilter(date_from=date(2026, 8, 1), date_to=date(2026, 8, 31)),
        )

    assert job.status == "succeeded"
    assert job.record_count == 1
    workbook = openpyxl.load_workbook(_workbook_path(job))
    sheet = workbook.active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert len(rows) == 3
    assert rows[2][0] == in_range.work_date.isoformat()


async def test_create_from_filter_with_no_matches_produces_header_only_workbook(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=None, filter_=ExportFilter()
        )

    assert job.status == "succeeded"
    assert job.record_count == 0
    workbook = openpyxl.load_workbook(_workbook_path(job))
    sheet = workbook.active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert rows == [(REPORT_TITLE,), (_DATE_HEADER,)]


async def test_create_marks_job_failed_without_raising_when_generation_fails(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    broken_settings = export_settings.model_copy(
        update={"export_temp_dir": export_settings.export_temp_dir / "missing" / "nested"}
    )
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)

    async with session_factory() as session:
        job = await ExportService(session, broken_settings).create(
            user.id, daily_report_day_ids=None, filter_=ExportFilter()
        )

    assert job.status == "failed"
    assert job.error_message is not None
    assert job.file_path is None
    assert job.record_count == 0


async def test_get_download_expires_lazily_and_deletes_the_file(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)

    async with session_factory() as session:
        job = await ExportService(session, export_settings, clock=_past_clock).create(
            user.id, daily_report_day_ids=None, filter_=ExportFilter()
        )
    file_path = _workbook_path(job)
    assert file_path.exists()

    async with session_factory() as session:
        with pytest.raises(ExportFileNotAvailableError):
            await ExportService(session, export_settings).get_download(user.id, job.id)

    assert not file_path.exists()
    async with session_factory() as session:
        expired = await ExportService(session, export_settings).get(user.id, job.id)
    assert expired.status == "expired"
    assert expired.file_path is None


async def test_get_download_rejects_files_outside_the_export_directory(
    export_engine: AsyncEngine, export_settings: Settings, tmp_path: Path
) -> None:
    outside_file = tmp_path / "outside.xlsx"
    outside_file.write_bytes(b"not a real workbook")
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=None, filter_=ExportFilter()
        )
        job.file_path = str(outside_file)
        await session.commit()

    async with session_factory() as session:
        with pytest.raises(ExportFileNotAvailableError):
            await ExportService(session, export_settings).get_download(user.id, job.id)
    assert outside_file.exists()


async def test_cleanup_expired_sweeps_every_owner_and_leaves_fresh_jobs_alone(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await _bootstrap_user(session)

    async with session_factory() as session:
        stale_job = await ExportService(session, export_settings, clock=_past_clock).create(
            user.id, daily_report_day_ids=None, filter_=ExportFilter()
        )
    async with session_factory() as session:
        fresh_job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=None, filter_=ExportFilter()
        )

    stale_path = _workbook_path(stale_job)
    fresh_path = _workbook_path(fresh_job)
    assert stale_path.exists()
    assert fresh_path.exists()

    async with session_factory() as session:
        await ExportService(session, export_settings).cleanup_expired()

    assert not stale_path.exists()
    assert fresh_path.exists()
    async with session_factory() as session:
        refreshed_stale = await ExportService(session, export_settings).get(user.id, stale_job.id)
        refreshed_fresh = await ExportService(session, export_settings).get(user.id, fresh_job.id)
    assert refreshed_stale.status == "expired"
    assert refreshed_fresh.status == "succeeded"


def test_export_media_type_is_the_standard_xlsx_mime_type() -> None:
    assert EXPORT_MEDIA_TYPE == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@pytest.mark.parametrize(
    ("username", "expected"),
    [
        ("alice", "alice"),
        ("张三", "张三"),
        ("ali/ce", "ali_ce"),
        ("../../etc/passwd", "etc_passwd"),
        ("  bob  ", "bob"),
        ("😀", "user"),
        ("", "user"),
    ],
)
def test_safe_filename_component_strips_characters_outside_the_electron_whitelist(
    username: str, expected: str
) -> None:
    """Must stay in lockstep with `SAFE_FILE_NAME_PATTERN` in

    `electron/src/main/export/file-saver.ts` — usernames have no character
    restriction beyond length, so this is the only thing standing between an
    admin-chosen username and an export the desktop app's save dialog
    whitelist would otherwise reject.
    """
    assert _safe_filename_component(username) == expected


async def test_create_sanitizes_an_unsafe_display_name_in_the_file_name(
    export_engine: AsyncEngine, export_settings: Settings
) -> None:
    session_factory = create_session_factory(export_engine)
    async with session_factory() as session:
        user = await BootstrapService(session).bootstrap(
            username="alice", password=STAGE5_PASSWORD, display_name="ali/ce bob"
        )
        archived = await _archived_day(
            session, owner_id=user.id, work_date=date(2026, 8, 5), fields=[], content={}
        )

    async with session_factory() as session:
        job = await ExportService(session, export_settings).create(
            user.id, daily_report_day_ids=[archived.id], filter_=None
        )

    assert job.file_name == "日报_ali_ce_bob_2026年8月5日.xlsx"
